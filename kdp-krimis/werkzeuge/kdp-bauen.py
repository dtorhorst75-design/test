#!/usr/bin/env python3
"""Baut einen Band der Reihe „Alles muss raus“ als KDP-fertige Word-Datei,
dazu die Bildprompts und ein Blatt mit allen KDP-Angaben.

Aufruf: python3 kdp-bauen.py ../band-01-starnberg

Textformat der Kapitel (text/01.txt …):
  # Kapiteltitel              erste Zeile; „# * Titel“ = Sonderseite ohne Kapitelnummer
  Absätze durch Leerzeilen getrennt
  ! 3                         Bildplatz Nr. 3 (aus band.json „bilder“)
  [[DOK Art | Überschrift]]   Dokument zum Mitraten, Art: Liste, Zettel, Brief, Notiz, Zeitung, Formular
  …                           bei Liste: Zeilen mit „ | “ getrennt, erste Zeile = Spaltenköpfe
  [[/DOK]]
"""
import importlib.util
import json
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HIER = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location('kb', os.path.join(HIER, '..', '..', 'krimis', 'werkzeuge', 'krimi-bauen.py'))
kb = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(kb)

GRAU = RGBColor(0x6B, 0x6B, 0x6B)
MONO = 'Courier New'
DOK = re.compile(r'\[\[DOK\s+(\w+)\s*\|\s*(.*?)\]\]\n(.*?)\n\[\[/DOK\]\]', re.S)


def lade_kapitel(ordner):
    kapitel = []
    for f in sorted(os.listdir(os.path.join(ordner, 'text'))):
        if not f.endswith('.txt'):
            continue
        roh = open(os.path.join(ordner, 'text', f), encoding='utf-8').read().strip()
        kopf, _, rest = roh.partition('\n')
        teile, pos = [], 0
        for m in DOK.finditer(rest):
            teile += [('text', a.strip()) for a in rest[pos:m.start()].split('\n\n') if a.strip()]
            teile.append(('dok', (m.group(1), m.group(2).strip(), m.group(3).strip())))
            pos = m.end()
        teile += [('text', a.strip()) for a in rest[pos:].split('\n\n') if a.strip()]
        titel = kopf.lstrip('#').strip()
        sonder = titel.startswith('*')
        kapitel.append((titel.lstrip('* ').strip(), sonder, teile))
    return kapitel


def zelle_schattieren(zelle, farbe):
    tcPr = zelle._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), farbe)
    tcPr.append(shd)


def tabellenrahmen(tabelle, art='single', staerke='6', farbe='7F7F7F'):
    tblPr = tabelle._tbl.tblPr
    bdr = OxmlElement('w:tblBorders')
    for seite in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{seite}')
        for k, v in (('val', art), ('sz', staerke), ('space', '0'), ('color', farbe)):
            el.set(qn(f'w:{k}'), v)
        bdr.append(el)
    tblPr.append(bdr)


def lauf(p, text, schrift=None, groesse=None, fett=False, kursiv=False, farbe=None):
    fett_an = strich_an = False
    for teil in re.split(r'(\*\*|~~)', text):
        if teil == '**':
            fett_an = not fett_an
            continue
        if teil == '~~':
            strich_an = not strich_an
            continue
        if not teil:
            continue
        r = p.add_run(teil)
        r.bold = fett or fett_an
        r.font.strike = strich_an
        r.italic = kursiv
        if schrift:
            r.font.name = schrift
            r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), schrift)
        if groesse:
            r.font.size = groesse
        if farbe:
            r.font.color.rgb = farbe


def dokument(doc, art, titel, inhalt):
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    if art == 'Liste':
        zeilen = [[z.strip() for z in l.split('|')] for l in inhalt.split('\n') if l.strip()]
        spalten = max(len(z) for z in zeilen)
        kopf = doc.add_paragraph()
        kopf.paragraph_format.space_after = Pt(3)
        kopf.paragraph_format.keep_with_next = True
        lauf(kopf, titel, MONO, Pt(8.5), fett=True)
        t = doc.add_table(rows=0, cols=spalten)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tabellenrahmen(t)
        # erste Spalte schmal, der Rest nach Textmenge verteilt (Satzspiegel 114 mm)
        laengen = [max(len(z[j]) if j < len(z) else 0 for z in zeilen) for j in range(spalten)]
        summe = sum(min(l, 60) + 8 for l in laengen[1:]) or 1
        breiten = [10] + [104 * (min(l, 60) + 8) / summe for l in laengen[1:]]
        for i, z in enumerate(zeilen):
            zellen = t.add_row().cells
            for j in range(spalten):
                zellen[j].width = Mm(breiten[j])
                p = zellen[j].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                lauf(p, z[j] if j < len(z) else '', MONO, Pt(7.5), fett=(i == 0))
                if i == 0:
                    zelle_schattieren(zellen[j], 'E7E7E7')
        for j, spalte in enumerate(t.columns):
            spalte.width = Mm(breiten[j])
        layout = OxmlElement('w:tblLayout')
        layout.set(qn('w:type'), 'fixed')
        t._tbl.tblPr.append(layout)
    else:
        t = doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabellenrahmen(t, 'dashed' if art == 'Zettel' else 'single', '8')
        zelle = t.rows[0].cells[0]
        zelle_schattieren(zelle, 'F4F4F4' if art != 'Zettel' else 'FBFBF6')
        schrift = MONO if art in ('Notiz', 'Formular') else None
        kursiv = art in ('Zettel', 'Brief')
        p = zelle.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        lauf(p, titel, MONO, Pt(8), fett=True, farbe=GRAU)
        for absatz_text in inhalt.split('\n\n'):
            p = zelle.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            if art == 'Zeitung':
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for k, zeile in enumerate(absatz_text.split('\n')):
                if k:
                    p.add_run().add_break()
                lauf(p, zeile, schrift, Pt(8.5 if schrift else 9.5), kursiv=kursiv)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def titelseite(doc, band):
    for _ in range(5):
        doc.add_paragraph()
    kb.absatz(doc, band['reihe'], einzug=False, groesse=Pt(13), farbe=GRAU, ausrichtung=WD_ALIGN_PARAGRAPH.CENTER, nach=4)
    kb.absatz(doc, band['reihenzusatz'], einzug=False, kursiv=True, groesse=Pt(10), farbe=GRAU, ausrichtung=WD_ALIGN_PARAGRAPH.CENTER, nach=30)
    kb.absatz(doc, band['titel'], einzug=False, groesse=Pt(24), ausrichtung=WD_ALIGN_PARAGRAPH.CENTER, nach=8)
    kb.absatz(doc, band['band'], einzug=False, kursiv=True, groesse=Pt(12), ausrichtung=WD_ALIGN_PARAGRAPH.CENTER, nach=50)
    kb.absatz(doc, band['autorin'], einzug=False, groesse=Pt(14), ausrichtung=WD_ALIGN_PARAGRAPH.CENTER)
    kb.seitenumbruch(doc)


def impressum(doc, band):
    for _ in range(14):
        doc.add_paragraph()
    for z in [
        f"**{band['reihe']} – {band['band']}: {band['titel']}**",
        '© 2026 Diana [Nachname]. Alle Rechte vorbehalten.',
        'Herausgeberin: Diana [Nachname], [Straße Hausnummer], [PLZ Ort]',
        'Kontakt: [E-Mail-Adresse]',
        'Druck: Amazon Distribution GmbH, Leipzig (bzw. siehe letzte Seite)',
        '',
        'Dieser Roman ist frei erfunden. Die Handlung spielt an realen Orten, alle Figuren, Firmen und Ereignisse sind jedoch erfunden. Ähnlichkeiten mit lebenden oder verstorbenen Personen sind zufällig und nicht beabsichtigt.',
        'Der Text wurde mit Unterstützung künstlicher Intelligenz erstellt und von der Autorin überarbeitet.',
    ]:
        kb.absatz(doc, z, einzug=False, groesse=Pt(8.5), nach=3)
    kb.seitenumbruch(doc)


def textseite(doc, titel, absaetze):
    kb.absatz(doc, titel, einzug=False, groesse=Pt(18), nach=14)
    for t in absaetze:
        kb.absatz(doc, t, einzug=False, nach=7)
    kb.seitenumbruch(doc)


def figurenseite(doc, figuren):
    kb.absatz(doc, 'Wer ist wer', einzug=False, groesse=Pt(18), nach=14)
    for name, wer in figuren:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.left_indent = Mm(4)
        p.paragraph_format.first_line_indent = Mm(-4)
        lauf(p, name, fett=True)
        lauf(p, ' – ' + wer)
    kb.seitenumbruch(doc)


def inhalt(doc, kapitel, seiten, extra):
    kb.absatz(doc, 'Inhalt', einzug=False, groesse=Pt(18), nach=14)
    n = 0
    for titel, sonder, _ in kapitel:
        if not sonder:
            n += 1
        anzeige = titel if sonder else f'{n}  {titel}'
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(114), alignment=2, leader=1)
        p.add_run(anzeige)
        p.add_run('\t' + str(seiten.get(titel, '')))
    for t in extra:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(114), alignment=2, leader=1)
        p.add_run(t)
        p.add_run('\t' + str(seiten.get(t, '')))
    kb.seitenumbruch(doc)


def buch(band, kapitel, seiten, ziel):
    doc = Document()
    kb.grundlayout(doc)
    titelseite(doc, band)
    impressum(doc, band)
    textseite(doc, 'So rätst du mit', band['mitraten'])
    figurenseite(doc, band['figuren'])
    extra = [band['nachwort_titel']]
    inhalt(doc, kapitel, seiten, extra)
    n = 0
    for titel, sonder, teile in kapitel:
        if sonder:
            for _ in range(3):
                doc.add_paragraph()
            kb.absatz(doc, titel, einzug=False, groesse=Pt(18), ausrichtung=WD_ALIGN_PARAGRAPH.CENTER, nach=20)
        else:
            n += 1
            kb.kapitelkopf(doc, n, titel)
        erster = True
        for art, wert in teile:
            if art == 'dok':
                dokument(doc, *wert)
                erster = True
            elif wert.startswith('! '):
                kb.bildplatz(doc, wert[2:].strip(), band['bilder'][wert[2:].strip()])
                erster = True
            else:
                kb.absatz(doc, wert, einzug=not (erster or sonder), nach=(6 if sonder else 0))
                erster = False
        kb.seitenumbruch(doc)
    kb.absatz(doc, band['nachwort_titel'], einzug=False, groesse=Pt(18), nach=14)
    for t in band['nachwort']:
        kb.absatz(doc, t, einzug=False, nach=8)
    kb.seitenumbruch(doc)
    for _ in range(4):
        doc.add_paragraph()
    for t in band['schluss']:
        kb.absatz(doc, t, einzug=False, ausrichtung=WD_ALIGN_PARAGRAPH.CENTER, nach=8)
    doc.save(ziel)


def bildprompts(band, kapitel, seiten, ziel):
    doc = Document()
    kb.grundlayout(doc)
    kb.absatz(doc, f"{band['reihe']} – {band['band']}: {band['titel']}", einzug=False, groesse=Pt(10), farbe=GRAU, nach=4)
    kb.absatz(doc, 'Bildprompts für Gemini', einzug=False, groesse=Pt(20), nach=12)
    for t in band['bildhinweise']:
        kb.absatz(doc, t, einzug=False, nach=6)
    kap = {}
    n = 0
    for titel, sonder, teile in kapitel:
        if not sonder:
            n += 1
        for art, wert in teile:
            if art == 'text' and wert.startswith('! '):
                kap[wert[2:].strip()] = (n, titel)
    for nr, bild in band['bilder'].items():
        k, titel = kap.get(nr, ('?', '?'))
        kb.absatz(doc, f'Bild {nr}: {bild["titel"]}', einzug=False, groesse=Pt(14), vor=14, nach=2)
        kb.absatz(doc, f'Kapitel {k} „{titel}“ · Seite {seiten.get("bild:" + nr, "?")}', einzug=False, farbe=GRAU, groesse=Pt(9), nach=4)
        kb.absatz(doc, bild['beschreibung'], einzug=False, nach=4)
        figuren = ' '.join(band['figurenprompt'][f] + '.' for f in bild.get('figuren', []))
        prompt = bild['prompt'] + (' Characters: ' + figuren if figuren else '') + '\n\n' + band['stil_sw']
        p = kb.absatz(doc, prompt, einzug=False, groesse=Pt(9), nach=6)
        kb.rahmen(p, 'BDBDBD')
    kb.absatz(doc, 'Farbversion (für das E-Book)', einzug=False, groesse=Pt(12), vor=14, nach=4)
    kb.absatz(doc, 'Ersetze den letzten Absatz eines Prompts durch diesen Stil:', einzug=False, nach=4)
    kb.absatz(doc, band['stil_farbe'], einzug=False, groesse=Pt(9), farbe=GRAU)
    doc.save(ziel)


def kdp_infos(band, worte, seitenzahl, ziel):
    doc = Document()
    kb.grundlayout(doc)
    kb.absatz(doc, f"{band['reihe']} – {band['band']}: {band['titel']}", einzug=False, groesse=Pt(10), farbe=GRAU, nach=4)
    kb.absatz(doc, 'Alles für Amazon KDP', einzug=False, groesse=Pt(20), nach=12)
    for ueberschrift, inhalt_ in band['kdp']:
        kb.absatz(doc, ueberschrift, einzug=False, groesse=Pt(13), vor=10, nach=4)
        for t in inhalt_:
            p = kb.absatz(doc, t, einzug=False, nach=5)
    kb.absatz(doc, 'Technische Daten dieses Manuskripts', einzug=False, groesse=Pt(13), vor=10, nach=4)
    kb.absatz(doc, f'Format A5 (148 × 210 mm, bei KDP „5,83 × 8,27 Zoll“), {seitenzahl} Seiten ohne Bilder, rund {worte} Wörter. '
                   'Die Ränder (17 mm) passen für KDP. Vor dem Hochladen in Word „Speichern unter › PDF“ wählen und das PDF hochladen.', einzug=False, nach=5)
    doc.save(ziel)


def main():
    ordner = os.path.abspath(sys.argv[1])
    band = json.load(open(os.path.join(ordner, 'band.json'), encoding='utf-8'))
    kapitel = lade_kapitel(ordner)
    ziel = os.path.join(ordner, band['datei'] + '.docx')
    buch(band, kapitel, {}, ziel)
    alle = [(t, None) for t, _, _ in kapitel]
    seiten = kb.seiten_ermitteln(ziel, alle, [band['nachwort_titel']])
    buch(band, kapitel, seiten, ziel)
    bildprompts(band, kapitel, seiten, os.path.join(ordner, band['datei'] + '-Bildprompts.docx'))
    worte = sum(len(re.sub(r'\*\*', '', w if a == 'text' else w[2]).split()) for _, _, t in kapitel for a, w in t)
    seitenzahl = max(v for v in seiten.values()) + 2 if seiten else 0
    kdp_infos(band, worte, seitenzahl, os.path.join(ordner, band['datei'] + '-KDP-Angaben.docx'))
    fehlend = [t for t, _, _ in kapitel if t not in seiten]
    print(f'{len(kapitel)} Abschnitte, {worte} Wörter, ca. {seitenzahl} Seiten')
    print('Seiten:', seiten)
    if fehlend:
        print('FEHLEND:', fehlend)


if __name__ == '__main__':
    main()
